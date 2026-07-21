import os
import openpyxl
from config import COLUMNS, SHEET_NAME, HEADER_MIGRATION
from core.excel.utils import is_file_locked, backup_file, get_header_map, row_to_dict, migrate_headers


class ExcelClient:
    """Handles CRUD operations on the inventory Excel file."""

    def __init__(self, file_path: str):
        self.file_path = file_path
        migrate_headers(file_path, HEADER_MIGRATION, SHEET_NAME)

    def _load_workbook(self):
        return openpyxl.load_workbook(self.file_path)

    def is_locked(self) -> bool:
        return is_file_locked(self.file_path)

    def get_all_items(self) -> list[dict]:
        wb    = self._load_workbook()
        ws    = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        h_map = get_header_map(ws)
        items = []
        for row_num in range(2, ws.max_row + 1):
            item = row_to_dict(ws, row_num, COLUMNS, h_map, check_empty=True)
            if item is not None:
                items.append(item)
        wb.close()
        return items

    def add_item(self, data: dict) -> int:
        wb    = self._load_workbook()
        ws    = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        h_map = get_header_map(ws)
        new_row = ws.max_row + 1
        for col_name in COLUMNS:
            col_idx = h_map.get(col_name.upper())
            if col_idx:
                ws.cell(row=new_row, column=col_idx, value=data.get(col_name, ""))
        backup_file(self.file_path)
        wb.save(self.file_path)
        wb.close()
        return new_row

    def update_item(self, row_num: int, data: dict) -> None:
        wb    = self._load_workbook()
        ws    = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        h_map = get_header_map(ws)
        for col_name in COLUMNS:
            col_idx = h_map.get(col_name.upper())
            if col_idx:
                ws.cell(row=row_num, column=col_idx, value=data.get(col_name, ""))
        backup_file(self.file_path)
        wb.save(self.file_path)
        wb.close()

    def delete_item(self, row_num: int) -> None:
        wb = self._load_workbook()
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        ws.delete_rows(row_num)
        backup_file(self.file_path)
        wb.save(self.file_path)
        wb.close()

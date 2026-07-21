import os
import sys
import logging

log = logging.getLogger("config")

_parent = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _parent not in sys.path:
    sys.path.insert(0, _parent)

from path_config import get_equipos_path, get_equipos_file_path

def resource_path(relative_path):
    base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

USERS_JSON_PATH = resource_path("users.json")

FILE_NAME = "Formato_Inventario_TI.xlsx"
SHEET_NAME = "Sheet1"

HEADER_MIGRATION = {
    "FIRST NAME": "Nombre",
    "LAST NAME": "Apellido",
    "LAPTOP/DESKTOP": "Tipo de Equipo",
    "DEVICE SERIAL": "No. Serie",
    "MONITOR 1": "Monitor 1",
    "MONITOR 1 SERIAL": "Serie Monitor 1",
    "MONITOR 2": "Monitor 2",
    "MONITOR 2 SERIAL": "Serie Monitor 2",
    "LOCATION": "Ubicacion",
}

BASE_COLUMNS = list(HEADER_MIGRATION.values())
COLUMNS = list(BASE_COLUMNS)

EXTRA_COLUMN_INDICES = list(range(9, 14))


def load_extra_columns(file_path: str) -> None:
    """Read column headers at indices 9–13 from the Excel file and extend COLUMNS."""
    global COLUMNS
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        headers = [cell.value for cell in ws[1]]
        wb.close()

        extra = []
        for idx in EXTRA_COLUMN_INDICES:
            if idx < len(headers) and headers[idx] is not None:
                name = str(headers[idx]).strip()
                if name:
                    extra.append(name.upper())

        COLUMNS.clear()
        COLUMNS.extend(BASE_COLUMNS + extra)
    except Exception as e:
        log.error("load_extra_columns failed: %s", e)
        COLUMNS.clear()
        COLUMNS.extend(BASE_COLUMNS)

_shared_dir_cache = None

def find_shared_dir() -> str | None:
    global _shared_dir_cache
    if _shared_dir_cache is not None:
        return _shared_dir_cache
    _shared_dir_cache = get_equipos_path()
    return _shared_dir_cache

def find_file_path() -> str | None:
    shared_dir = find_shared_dir()
    if shared_dir:
        candidate = os.path.join(shared_dir, FILE_NAME)
        if os.path.exists(candidate):
            return candidate
    return None

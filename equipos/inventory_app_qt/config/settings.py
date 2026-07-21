import os
import sys
import glob
import logging

log = logging.getLogger("equipos_settings")

_parent_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
if _parent_root not in sys.path:
    sys.path.insert(0, _parent_root)

from path_config import get_equipos_path, get_equipos_file_path

APP_NAME = "IT Inventory - Equipos"
APP_VERSION = "1.0.0"

FILE_NAME = "Formato_Inventario_TI.xlsx"
SHEET_NAME = "Sheet1"

HEADER_MIGRATION = {
    "FIRST NAME": "Nombre",
    "LAST NAME": "Apellido",
    "LAPTOP/DESKTOP": "Tipo de Equipo",
    "DEVICE SERIAL": "No. Serie",
    "MONITOR 1": "Monitor 1",
    "MONITOR 1 SERIAL": "Serie Monitor 1",
    "MONITOR 2": "Monitor 2",
    "MONITOR 2 SERIAL": "Serie Monitor 2",
    "LOCATION": "Ubicacion",
}

BASE_COLUMNS = list(HEADER_MIGRATION.values())

EXTRA_COLUMN_INDICES = list(range(9, 14))


def resource_path(relative_path):
    # Bundled files are in equipos/ root
    return os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), relative_path)


USERS_JSON_PATH = resource_path("users.json")


def get_onedrive_base_path():
    path = get_equipos_path()
    if path:
        log.info("OneDrive base path resolved: %s", path)
        return path
    username = os.environ.get("USERNAME") or os.environ.get("USER") or "unknown"
    fallback = rf"C:\Users\{username}\OneDrive - a360inc\PTY Files - EQUIPOS"
    log.warning("get_equipos_path() returned None — falling back: %s", fallback)
    return fallback


def get_excel_path():
    path = get_equipos_file_path()
    if path:
        return path
    return os.path.join(get_onedrive_base_path(), FILE_NAME)


def load_extra_columns(file_path):
    global COLUMNS
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        headers = [cell.value for cell in ws[1]]
        wb.close()

        extra = []
        for idx in EXTRA_COLUMN_INDICES:
            if idx < len(headers) and headers[idx] is not None:
                name = str(headers[idx]).strip()
                if name:
                    extra.append(name.upper())

        COLUMNS.clear()
        COLUMNS.extend(BASE_COLUMNS + extra)
    except Exception as e:
        log.error("load_extra_columns failed: %s", e)
        COLUMNS.clear()
        COLUMNS.extend(BASE_COLUMNS)


COLUMNS = BASE_COLUMNS.copy()

USERS_DIR = get_onedrive_base_path()
USERS_FILE = os.path.join(USERS_DIR, "users.json")
log.info("USERS_FILE: %s", USERS_FILE)

ROLES = {
    "admin": "Admin",
    "editor": "Editor",
    "viewer": "Viewer",
}

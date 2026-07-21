import os
import openpyxl

from config.settings import COLUMNS, SHEET_NAME, get_excel_path
from core.excel.utils import is_file_locked, backup_file, get_header_map, row_to_dict


class ExcelService:

    def __init__(self):
        self._filepath = None
        self._wb = None

    @property
    def filepath(self):
        return get_excel_path()

    def _load_workbook(self):
        path = self.filepath
        if not os.path.exists(path):
            raise FileNotFoundError(f"Excel file not found: {path}")
        if is_file_locked(path):
            raise PermissionError(
                "The Excel file is currently open in another program.\n"
                "Please close it and try again."
            )
        self._wb = openpyxl.load_workbook(path)
        return self._wb

    def _get_sheet(self):
        if SHEET_NAME in self._wb.sheetnames:
            return self._wb[SHEET_NAME]
        return self._wb.active

    def _save(self):
        self._wb.save(self.filepath)

    def get_all_items(self):
        self._load_workbook()
        ws = self._get_sheet()
        h_map = get_header_map(ws)
        items = []
        for row_num in range(2, ws.max_row + 1):
            item = row_to_dict(ws, row_num, COLUMNS, h_map, check_empty=True)
            if item is not None:
                items.append(item)
        return items

    def add_item(self, data):
        self._load_workbook()
        ws = self._get_sheet()
        h_map = get_header_map(ws)
        new_row = ws.max_row + 1
        for col_name in COLUMNS:
            col_idx = h_map.get(col_name.upper())
            if col_idx:
                ws.cell(row=new_row, column=col_idx, value=data.get(col_name, ""))
        backup_file(self.filepath)
        self._save()
        return new_row

    def update_item(self, row_num, data):
        self._load_workbook()
        ws = self._get_sheet()
        h_map = get_header_map(ws)
        for col_name in COLUMNS:
            col_idx = h_map.get(col_name.upper())
            if col_idx:
                ws.cell(row=row_num, column=col_idx, value=data.get(col_name, ""))
        backup_file(self.filepath)
        self._save()

    def delete_item(self, row_num):
        self._load_workbook()
        ws = self._get_sheet()
        ws.delete_rows(row_num)
        backup_file(self.filepath)
        self._save()

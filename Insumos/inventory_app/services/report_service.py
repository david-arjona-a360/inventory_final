import os
from datetime import datetime

from fpdf import FPDF
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADER_FILL = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
HEADER_FONT = Font(bold=True, color="000000", size=11)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


class ReportService:

    def __init__(self, excel_service):
        self._excel = excel_service

    def generate_report_data(self, filters=None):
        items = self._excel.get_all_items(include_inactive=True)
        if not filters:
            return items
        filtered = []
        for item in items:
            match = True
            if filters.get("categoria"):
                if item.get("categoria", "").lower() != filters["categoria"].lower():
                    match = False
            if filters.get("producto"):
                if filters["producto"].lower() not in item.get("producto", "").lower():
                    match = False
            if filters.get("cantidad_min") is not None:
                try:
                    if int(item.get("cantidad", 0)) < int(filters["cantidad_min"]):
                        match = False
                except (ValueError, TypeError):
                    match = False
            if filters.get("cantidad_max") is not None:
                try:
                    if int(item.get("cantidad", 0)) > int(filters["cantidad_max"]):
                        match = False
                except (ValueError, TypeError):
                    match = False
            if match:
                filtered.append(item)
        return filtered

    def export_to_excel(self, data, filepath):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte"
        headers = ["Producto", "Marca", "Tamaño", "Cantidad", "Categoria", "Proveedor", "Cant. Minima"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        for row_idx, item in enumerate(data, 2):
            values = [
                item.get("producto", ""),
                item.get("marca", ""),
                item.get("tamano", ""),
                item.get("cantidad", 0),
                item.get("categoria", ""),
                item.get("proveedor", ""),
                item.get("cant_minima", 0),
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = THIN_BORDER
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 14
        wb.save(filepath)

    def export_to_pdf(self, data, filepath):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Reporte de Inventario - Insumos", ln=True, align="C")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 8)
        pdf.cell(0, 5, f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True, align="R")
        pdf.ln(5)
        headers = ["Producto", "Marca", "Tamaño", "Cantidad", "Categoria", "Proveedor", "Cant. Minima"]
        col_widths = [45, 30, 30, 18, 28, 30, 18]
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(204, 204, 204)
        pdf.set_text_color(0, 0, 0)
        for i, header in enumerate(headers):
            pdf.cell(col_widths[i], 7, header, border=1, align="C", fill=True)
        pdf.ln()
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(0, 0, 0)
        for item in data:
            values = [
                item.get("producto", ""),
                item.get("marca", ""),
                item.get("tamano", ""),
                str(item.get("cantidad", 0)),
                item.get("categoria", ""),
                item.get("proveedor", ""),
                str(item.get("cant_minima", 0)),
            ]
            max_h = 7
            for i, val in enumerate(values):
                pdf.cell(col_widths[i], max_h, val[:40], border=1, align="C")
            pdf.ln()
        pdf.output(filepath)

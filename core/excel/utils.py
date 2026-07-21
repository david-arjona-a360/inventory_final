import os
import shutil
import tempfile


def is_file_locked(filepath):
    if not os.path.exists(filepath):
        return False
    folder = os.path.dirname(filepath)
    filename = os.path.basename(filepath)
    lock = os.path.join(folder, f"~${filename}")
    return os.path.exists(lock)


def safe_copy_for_reading(filepath):
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    shutil.copy2(filepath, tmp.name)
    return tmp.name


def backup_file(filepath):
    backup_path = filepath + ".bak"
    try:
        shutil.copy2(filepath, backup_path)
    except OSError:
        pass


def get_header_map(sheet):
    return {
        cell.value.strip().upper(): cell.column
        for cell in sheet[1]
        if cell.value
    }


def row_to_dict(sheet, row_num, columns, header_map, check_empty=False):
    item = {"_row": row_num}
    all_empty = True
    for col_name in columns:
        col_idx = header_map.get(col_name.upper())
        val = sheet.cell(row=row_num, column=col_idx).value or "" if col_idx else ""
        item[col_name] = val
        if check_empty and str(val).strip():
            all_empty = False
    if check_empty and all_empty:
        return None
    return item


def migrate_headers(filepath, header_map, sheet_name="Sheet1"):
    import openpyxl
    if not os.path.exists(filepath) or is_file_locked(filepath):
        return False
    wb = openpyxl.load_workbook(filepath)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]
    changed = False
    for col_idx, header in enumerate(headers, 1):
        if header and header in header_map:
            ws.cell(row=1, column=col_idx, value=header_map[header])
            changed = True
    if changed:
        backup_file(filepath)
        wb.save(filepath)
    wb.close()
    return changed

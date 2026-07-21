import os
import sys
import uuid
import logging
import tempfile
from datetime import datetime
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import (
    get_excel_path, MAIN_SHEET, HISTORY_SHEET, HIDDEN_COLUMNS,
    VISIBLE_COLUMNS, COLUMN_MAP
)
from core.excel.utils import is_file_locked

_DIAG_LOG_EXCEL = os.path.join(tempfile.gettempdir(), "inventory_diag.log")
logging.basicConfig(
    level=logging.DEBUG,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.FileHandler(_DIAG_LOG_EXCEL, mode="a"),
        logging.StreamHandler(sys.stderr),
    ],
    force=True,
)
log = logging.getLogger("excel_service")

HEADER_FILL = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
HEADER_FONT = Font(bold=True, color="000000", size=11)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


class ExcelService:

    def __init__(self):
        self._filepath = None
        self._wb = None

    @property
    def filepath(self):
        return get_excel_path()

    def _load_workbook(self):
        path = self.filepath
        log.info("_load_workbook: path=%s exists=%s", path, os.path.isfile(path))
        if is_file_locked(path):
            raise PermissionError(
                "The Excel file is currently open in another program.\n"
                "Please close it and try again."
            )
        if not os.path.exists(path):
            log.warning("_load_workbook: file NOT found at %s -> creating new workbook", path)
            self._create_workbook()
        else:
            log.info("_load_workbook: loading existing file: %s", path)
            self._wb = openpyxl.load_workbook(path)
            log.info("_load_workbook: sheets=%s", self._wb.sheetnames)
            if MAIN_SHEET in self._wb.sheetnames:
                ws = self._wb[MAIN_SHEET]
                log.info("_load_workbook: MAIN_SHEET '%s' rows=%d cols=%d", MAIN_SHEET, ws.max_row, ws.max_column)

    def _create_workbook(self):
        self._wb = openpyxl.Workbook()
        self._initialize_main_sheet()
        self._initialize_history_sheet()
        self._save()

    def _initialize_main_sheet(self):
        ws = self._wb.active
        ws.title = MAIN_SHEET
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        for col_idx, col_name in enumerate(all_columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        self._hide_columns(ws)

    def _hide_columns(self, ws):
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        for col_name in HIDDEN_COLUMNS:
            idx = all_columns.index(col_name) + 1
            ws.column_dimensions[get_column_letter(idx)].hidden = True

    def _initialize_history_sheet(self):
        if HISTORY_SHEET in self._wb.sheetnames:
            return
        ws = self._wb.create_sheet(title=HISTORY_SHEET)
        headers = ["item_id", "previous_quantity", "new_quantity", "user", "timestamp", "action"]
        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        ws.sheet_state = 'veryHidden'

    def _save(self):
        path = self.filepath
        if is_file_locked(path):
            raise PermissionError(
                "The Excel file is currently open in another program.\n"
                "Please close it and try again."
            )
        self._wb.save(path)

    def get_all_items(self, include_inactive=False):
        self._load_workbook()
        ws = self._wb[MAIN_SHEET]
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        items = []
        row_count = ws.max_row
        log.info("get_all_items: max_row=%d max_col=%d include_inactive=%s", row_count, ws.max_column, include_inactive)
        for row in range(2, row_count + 1):
            item = {}
            for col_idx, col_name in enumerate(all_columns, 1):
                val = ws.cell(row=row, column=col_idx).value
                item[COLUMN_MAP.get(col_name, col_name)] = val
            if not include_inactive and item.get("is_active") == 0:
                log.debug("get_all_items: row %d filtered out (inactive)", row)
                continue
            if item.get("id") is None:
                log.debug("get_all_items: row %d filtered out (id is None). cols=%s", row, {k: v for k, v in item.items() if v is not None})
                continue
            items.append(item)
        log.info("get_all_items: returning %d items", len(items))
        return items

    def get_item_by_id(self, item_id):
        self._load_workbook()
        ws = self._wb[MAIN_SHEET]
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        for row in range(2, ws.max_row + 1):
            cell_id = ws.cell(row=row, column=all_columns.index("id") + 1).value
            if str(cell_id) == str(item_id):
                item = {}
                for col_idx, col_name in enumerate(all_columns, 1):
                    val = ws.cell(row=row, column=col_idx).value
                    item[COLUMN_MAP.get(col_name, col_name)] = val
                return item, row
        return None, None

    def add_item(self, data, username):
        self._load_workbook()
        ws = self._wb[MAIN_SHEET]
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        new_id = str(uuid.uuid4())
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        row_data = {
            "producto": data.get("producto", ""),
            "marca": data.get("marca", ""),
            "tamano": data.get("tamano", ""),
            "cantidad": int(data.get("cantidad", 0)),
            "categoria": data.get("categoria", ""),
            "proveedor": data.get("proveedor", ""),
            "cant_minima": int(data.get("cant_minima", 0)) if data.get("cant_minima") else 0,
            "id": new_id,
            "created_at": now,
            "updated_at": now,
            "updated_by": username,
            "is_active": 1,
        }
        next_row = ws.max_row + 1
        for col_idx, col_name in enumerate(all_columns, 1):
            key = COLUMN_MAP.get(col_name, col_name)
            ws.cell(row=next_row, column=col_idx, value=row_data.get(key, ""))
        self._save()
        self._log_history(new_id, None, row_data["cantidad"], username, "create")
        return new_id

    def update_item(self, item_id, data, username, force=False):
        self._load_workbook()
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        item, row = self._find_item_by_id(item_id)
        if item is None:
            raise ValueError(f"Item with id {item_id} not found")

        ws = self._wb[MAIN_SHEET]

        if not force:
            stored_updated_at = item.get("updated_at", "")
            provided_updated_at = data.get("_original_updated_at", "")
            if stored_updated_at and provided_updated_at and stored_updated_at != provided_updated_at:
                return False, "conflict"

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        prev_qty = item.get("cantidad", 0)
        new_qty = int(data.get("cantidad", prev_qty))

        row_data = {
            "producto": data.get("producto", item["producto"]),
            "marca": data.get("marca", item["marca"]),
            "tamano": data.get("tamano", item["tamano"]),
            "cantidad": new_qty,
            "categoria": data.get("categoria", item["categoria"]),
            "proveedor": data.get("proveedor", item.get("proveedor", "")),
            "cant_minima": int(data.get("cant_minima", item.get("cant_minima", 0))),
            "id": item_id,
            "created_at": item.get("created_at", now),
            "updated_at": now,
            "updated_by": username,
            "is_active": int(data.get("is_active", item.get("is_active", 1))),
        }
        for col_idx, col_name in enumerate(all_columns, 1):
            key = COLUMN_MAP.get(col_name, col_name)
            ws.cell(row=row, column=col_idx, value=row_data.get(key, ""))
        self._save()
        self._log_history(item_id, prev_qty, new_qty, username, "update")
        return True, "ok"

    def _find_item_by_id(self, item_id):
        ws = self._wb[MAIN_SHEET]
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        for row in range(2, ws.max_row + 1):
            cell_id = ws.cell(row=row, column=all_columns.index("id") + 1).value
            if str(cell_id) == str(item_id):
                item = {}
                for col_idx, col_name in enumerate(all_columns, 1):
                    val = ws.cell(row=row, column=col_idx).value
                    item[COLUMN_MAP.get(col_name, col_name)] = val
                return item, row
        return None, None

    def soft_delete(self, item_id, username):
        return self.update_item(item_id, {"is_active": 0}, username)

    def restore_item(self, item_id, username):
        return self.update_item(item_id, {"is_active": 1}, username, force=True)

    def _log_history(self, item_id, prev_qty, new_qty, user, action):
        if HISTORY_SHEET not in self._wb.sheetnames:
            self._initialize_history_sheet()
        ws = self._wb[HISTORY_SHEET]
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        next_row = ws.max_row + 1
        ws.cell(row=next_row, column=1, value=str(item_id))
        ws.cell(row=next_row, column=2, value=prev_qty)
        ws.cell(row=next_row, column=3, value=new_qty)
        ws.cell(row=next_row, column=4, value=user)
        ws.cell(row=next_row, column=5, value=now)
        ws.cell(row=next_row, column=6, value=action)
        self._wb.save(self.filepath)

    def get_history(self):
        self._load_workbook()
        if HISTORY_SHEET not in self._wb.sheetnames:
            return []
        ws = self._wb[HISTORY_SHEET]
        rows = []
        for row in range(2, ws.max_row + 1):
            rows.append({
                "item_id": ws.cell(row=row, column=1).value,
                "previous_quantity": ws.cell(row=row, column=2).value,
                "new_quantity": ws.cell(row=row, column=3).value,
                "user": ws.cell(row=row, column=4).value,
                "timestamp": ws.cell(row=row, column=5).value,
                "action": ws.cell(row=row, column=6).value,
            })
        return rows

    def _insert_column_after(self, ws, after_col_name, new_col_name):
        all_columns = VISIBLE_COLUMNS + HIDDEN_COLUMNS
        existing = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val:
                existing[val] = c

        if new_col_name in existing:
            return False

        if after_col_name not in existing:
            col_idx = len(existing) + 1
        else:
            col_idx = existing[after_col_name] + 1
            ws.insert_cols(col_idx)

        cell = ws.cell(row=1, column=col_idx, value=new_col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER

        if new_col_name in HIDDEN_COLUMNS:
            for row in range(2, ws.max_row + 1):
                if new_col_name == "id":
                    ws.cell(row=row, column=col_idx, value=str(uuid.uuid4()))
                elif new_col_name == "created_at":
                    ws.cell(row=row, column=col_idx, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                elif new_col_name == "updated_at":
                    ws.cell(row=row, column=col_idx, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                elif new_col_name == "is_active":
                    ws.cell(row=row, column=col_idx, value=1)
                elif new_col_name == "updated_by":
                    ws.cell(row=row, column=col_idx, value="system")
        else:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx, value="")

        return True

    def ensure_structure(self):
        log.info("ensure_structure: filepath=%s", self.filepath)
        if not os.path.exists(self.filepath):
            log.warning("ensure_structure: file NOT found -> creating new workbook")
            self._create_workbook()
            return
        try:
            self._load_workbook()
        except PermissionError:
            raise
        modified = False
        ws = self._wb[MAIN_SHEET]

        existing = {ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)}
        log.info("ensure_structure: existing headers: %s", existing)

        if "Proveedor" not in existing:
            modified |= self._insert_column_after(ws, "Categoria", "Proveedor")
        if "Cant. Minima" not in existing:
            modified |= self._insert_column_after(ws, "Proveedor", "Cant. Minima")

        self._hide_columns(ws)
        if HISTORY_SHEET not in self._wb.sheetnames:
            self._initialize_history_sheet()
            modified = True
        if modified:
            log.info("ensure_structure: workbook modified, saving")
            self._save()
        log.info("ensure_structure: completed successfully")
        return True
